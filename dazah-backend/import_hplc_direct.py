"""液相色谱对照品Excel数据导入脚本 - 使用psycopg2直连"""

import openpyxl
from datetime import datetime, date
import psycopg2

# Database config
DATABASE_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "database": "dahzah",
    "user": "postgres",
    "password": "postgres"
}


def parse_date(val):
    """解析日期"""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        val = val.strip()
        for fmt in ["%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"]:
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                continue
        return None
    return None


def parse_number(val):
    """解析数字"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, str):
        val = val.strip()
        if not val or val == '/':
            return None
        try:
            return float(val) if '.' in val else int(val)
        except ValueError:
            return None
    return None


def parse_bool(val):
    """解析布尔值"""
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    if isinstance(val, str):
        return val.strip().upper() in ("是", "YES", "Y", "TRUE", "1", "有")
    return False


def generate_ref_code(name: str, idx: int) -> str:
    """生成对照品编号"""
    return f"REF{str(idx).zfill(4)}"


def import_hplc_reference():
    """导入液相色谱对照品数据"""
    excel_path = r"D:\Workspace\液相对照品台账.xlsx"

    print(f"读取Excel文件: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 2:
        print("Excel中没有数据行")
        return

    # 读取表头
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    header_to_idx = {h: i for i, h in enumerate(headers) if h}

    print(f"Headers count: {len(headers)}")
    print(f"Data rows: {len(rows) - 1}")

    # 定义字段映射
    FIELD_MAP = {
        "名称": "ref_name",
        "交接单": "handover_no",
        "项目": "project_name",
        "是否有COA": "has_coa",
        "开瓶日期": "open_date",
        "开瓶有效期": "open_expire_days",
        "现有库存": "stock_status",
        "厂内批号": "internal_batch",
        "CAS NO": "cas_no",
        "货号CAT NO": "cat_no",
        "到货日期": "arrival_date",
        "存放位置": "location",
        "规格/瓶": "spec",
        "纯度": "purity",
        "含量": "content",
        "数量": "quantity",
        "生产/标定日期": "produce_date",
        "复标期": "recal_cycle_days",
        "有效期": "expire_date",
        "储存条件": "storage_cond_code",
        "来源": "manufacturer",
    }

    conn = psycopg2.connect(**DATABASE_CONFIG)
    cursor = conn.cursor()

    success = 0
    failed = 0
    errors = []

    # 找编号列 (在厂内批号后面)
    ref_code_idx = None
    for i, h in enumerate(headers):
        if h == "编号" and i > 8:
            ref_code_idx = i
            break

    print(f"Using ref_code column index: {ref_code_idx}")

    for idx, row in enumerate(rows[1:], start=2):
        if not row or all(v is None for v in row):
            continue
        row_list = list(row)
        data = {}

        # 获取编号
        if ref_code_idx is not None and ref_code_idx < len(row_list):
            ref_code_val = row_list[ref_code_idx]
            if ref_code_val is not None and str(ref_code_val).strip():
                data["ref_code"] = str(ref_code_val).strip()

        # 处理其他字段
        for header, field_name in FIELD_MAP.items():
            col_idx = header_to_idx.get(header, -1)
            if col_idx < 0 or col_idx >= len(row_list):
                continue
            val = row_list[col_idx]
            if val is None or (isinstance(val, str) and not val.strip()):
                continue

            # 类型转换
            if field_name in ("open_date", "arrival_date", "produce_date", "expire_date"):
                val = parse_date(val)
            elif field_name in ("purity", "content", "quantity", "recal_cycle_days", "open_expire_days"):
                val = parse_number(val)
            elif field_name == "has_coa":
                val = parse_bool(val)
            elif isinstance(val, str):
                val = val.strip().replace('\n', '/')

            if val is not None:
                data[field_name] = val

        # 如果没有名称，跳过
        if not data.get("ref_name"):
            failed += 1
            errors.append(f"第{idx}行: 缺少对照品名称")
            continue

        # 如果没有编号，自动生成
        if not data.get("ref_code"):
            data["ref_code"] = generate_ref_code(data["ref_name"], idx)

        # 状态默认值
        data["ref_status"] = 0
        data["create_by"] = 1

        # 构建SQL
        columns = list(data.keys())
        values = list(data.values())
        placeholders = ','.join(['%s' for _ in values])

        sql = f"INSERT INTO public.t_qs_hplc_reference ({','.join(columns)}) VALUES ({placeholders})"

        try:
            cursor.execute(sql, values)
            success += 1
        except Exception as e:
            failed += 1
            err_msg = str(e)[:150]
            errors.append(f"第{idx}行 [{data.get('ref_code', 'N/A')}]: {err_msg}")

        if idx % 50 == 0:
            conn.commit()
            print(f"进度: 已处理 {idx} 行")

    conn.commit()
    cursor.close()
    conn.close()

    print(f"\n导入完成:")
    print(f"  成功: {success} 条")
    print(f"  失败: {failed} 条")
    if errors:
        print(f"\n前10条错误:")
        for err in errors[:10]:
            print(f"  {err}")


if __name__ == "__main__":
    import_hplc_reference()